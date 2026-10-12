# ============================================================
# Arma la HOJA DE ETIQUETADO para medir precision de tema y tono
# (PS-06 y PS-07). Toma una muestra de ~50 registros YA clasificados
# de los contenedores silver-* y los vuelca a un Excel con columnas
# vacias para que marques a mano el tema y el tono correctos.
#
# NO calcula nada. El calculo de precision se hace aparte, cuando
# devuelvas el Excel lleno.
# ============================================================

import os
import sys
import json
import random
from azure.storage.blob import BlobServiceClient
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PREFIJO = "silver-"
XLSX_SALIDA = "hoja_etiquetado_tema_tono.xlsx"
MUESTRA = 50
SEMILLA = 42  # reproducible

CAMPO_TEMA = "clasificacion.tipo_senal"
CAMPOS_TONO = ["clasificacion.tono", "clasificacion.sentimiento", "sentimiento"]
CAMPOS_TEXTO = ["titulo", "title", "resumen", "summary", "contenido",
                "content", "texto", "text", "descripcion", "cuerpo"]


def obtener_conexion():
    for k in ("AZURE_STORAGE_CONNECTION_STRING", "AzureWebJobsStorage"):
        v = os.environ.get(k)
        if v:
            return v
    d = os.getcwd()
    for _ in range(4):
        p = os.path.join(d, "local.settings.json")
        if os.path.exists(p):
            with open(p, encoding="utf-8") as f:
                vals = json.load(f).get("Values", {})
            if vals.get("AzureWebJobsStorage"):
                return vals["AzureWebJobsStorage"]
        d = os.path.dirname(d)
    return None


def valor(dic, ruta):
    cur = dic
    for parte in ruta.split("."):
        if isinstance(cur, dict) and parte in cur:
            cur = cur[parte]
        else:
            return None
    return cur


def primer_valor(reg, rutas):
    for r in rutas:
        v = valor(reg, r)
        if v is not None and str(v).strip() != "":
            return v
    return None


def texto_registro(reg):
    partes = []
    for c in CAMPOS_TEXTO:
        v = valor(reg, c)
        if isinstance(v, str) and v.strip():
            partes.append(v.strip())
    if not partes:  # respaldo: las cadenas mas largas del registro
        cadenas = sorted([v for v in reg.values() if isinstance(v, str)],
                         key=len, reverse=True)
        partes = cadenas[:2]
    t = " | ".join(partes)
    return t[:500]


def registros_de_blob(texto):
    texto = texto.strip()
    if not texto:
        return []
    try:
        obj = json.loads(texto)
        return obj if isinstance(obj, list) else [obj]
    except json.JSONDecodeError:
        regs = []
        for linea in texto.splitlines():
            if linea.strip():
                try:
                    regs.append(json.loads(linea))
                except json.JSONDecodeError:
                    pass
        return regs


def main():
    conn = obtener_conexion()
    if not conn:
        print("ERROR: no encontre la cadena de conexion.")
        print('  $env:AZURE_STORAGE_CONNECTION_STRING = "<cadena de stinteligencia>"')
        sys.exit(1)

    svc = BlobServiceClient.from_connection_string(conn)
    contenedores = [c.name for c in svc.list_containers(name_starts_with=PREFIJO)]
    if not contenedores:
        print("No encontre contenedores con prefijo '%s'." % PREFIJO)
        sys.exit(1)

    clasificados = []
    estructura = False
    for cont in contenedores:
        cc = svc.get_container_client(cont)
        for blob in cc.list_blobs():
            txt = cc.download_blob(blob.name).readall().decode("utf-8", errors="replace")
            for reg in registros_de_blob(txt):
                if not isinstance(reg, dict):
                    continue
                if not estructura:
                    print("Estructura del primer registro:")
                    for k, v in reg.items():
                        print("  %s: %s" % (k, ", ".join(v.keys()) if isinstance(v, dict) else type(v).__name__))
                    print()
                    estructura = True
                tema = valor(reg, CAMPO_TEMA)
                if tema is None or str(tema).strip() == "":
                    continue  # solo registros ya clasificados
                clasificados.append({
                    "contenedor": cont,
                    "id": reg.get("item_id_origen") or reg.get("id") or blob.name,
                    "texto": texto_registro(reg),
                    "tema_sistema": tema,
                    "tono_sistema": primer_valor(reg, CAMPOS_TONO) or "",
                })

    if not clasificados:
        print("No encontre registros con tema asignado (clasificacion.tipo_senal).")
        sys.exit(1)

    random.seed(SEMILLA)
    random.shuffle(clasificados)
    muestra = clasificados[:MUESTRA]
    print("Registros clasificados: %d  ->  muestra tomada: %d" % (len(clasificados), len(muestra)))

    temas = sorted({str(r["tema_sistema"]) for r in clasificados})
    tonos = sorted({str(r["tono_sistema"]) for r in clasificados if r["tono_sistema"]})
    if not tonos:
        tonos = ["positivo", "negativo", "neutral"]

    wb = Workbook()
    ws = wb.active
    ws.title = "etiquetado"
    cols = ["n", "contenedor", "id", "texto",
            "tema_correcto", "tono_correcto", "tema_sistema", "tono_sistema"]
    ws.append(cols)
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3864")
    for i, _ in enumerate(cols, 1):
        c = ws.cell(row=1, column=i)
        c.font = hdr
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, r in enumerate(muestra, 1):
        ws.append([i, r["contenedor"], str(r["id"]), r["texto"],
                   "", "", str(r["tema_sistema"]), str(r["tono_sistema"])])

    anchos = {"A": 5, "B": 22, "C": 26, "D": 70, "E": 18, "F": 16, "G": 18, "H": 16}
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    # listas desplegables (si caben en el limite de Excel)
    def agregar_lista(col_letra, opciones):
        texto = ",".join(opciones)
        if len(texto) <= 250:
            dv = DataValidation(type="list", formula1='"%s"' % texto, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add("%s2:%s%d" % (col_letra, col_letra, len(muestra) + 1))
    agregar_lista("E", temas)
    agregar_lista("F", tonos)

    # hoja de instrucciones
    ins = wb.create_sheet("instrucciones")
    pasos = [
        "COMO ETIQUETAR (PS-06 tema y PS-07 tono)",
        "",
        "1. Lee la columna 'texto' de cada fila.",
        "2. En 'tema_correcto' escribe el tema que TU consideras correcto.",
        "3. En 'tono_correcto' escribe el tono correcto (positivo / negativo / neutral).",
        "4. Hazlo con tu propio criterio ANTES de mirar las columnas del sistema",
        "   (tema_sistema, tono_sistema) para no sesgar tu juicio.",
        "5. Guarda el archivo y subelo para calcular la precision real.",
    ]
    for p in pasos:
        ins.append([p])
    ins.column_dimensions["A"].width = 80

    wb.save(XLSX_SALIDA)
    print("Archivo generado: %s" % os.path.abspath(XLSX_SALIDA))


if __name__ == "__main__":
    main()